"""Tests for XLSX export functionality."""

from collections.abc import Iterator
from io import BytesIO
from pathlib import Path
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from apps.api.dependencies import get_db_session
from apps.api.main import app
from apps.api.routers import documents
from apps.worker.processor import DocumentProcessor
from apps.worker.worker import Worker
from packages.common.db import Base
from packages.common.export import generate_leads_xlsx
from packages.common.queue import InMemoryQueue


class FakeStorage:
    def __init__(self, root: Path) -> None:
        self.root = root

    def put(self, *, job_id, document_id, filename, content) -> str:
        path = self.root / str(document_id)
        path.write_bytes(content)
        return f"local://{path}"


class FakeObjectReader:
    def read(self, *, source_uri: str) -> bytes:
        path = Path(source_uri.removeprefix("local://"))
        return path.read_bytes()


class FakeInferenceClient:
    model_name = "qwen/test"
    model_version = "test"

    def __init__(self) -> None:
        self._call_count = 0

    def extract(self, *, image_bytes, mime_type) -> dict:
        self._call_count += 1
        return {
            "first_name": f"First{self._call_count}",
            "last_name": f"Last{self._call_count}",
            "job_title": f"Title{self._call_count}",
            "company": f"Company{self._call_count}",
            "location": f"City{self._call_count}",
            "phone_number": f"+1{self._call_count:010d}",
            "email_address": f"user{self._call_count}@test.com",
        }


@pytest.fixture()
def isolated_engine(tmp_path):
    engine = create_engine(
        f"sqlite:///{tmp_path / 'test.db'}", future=True
    )
    from packages.common import models  # noqa: F401

    Base.metadata.create_all(engine)
    yield engine
    engine.dispose()


@pytest.fixture()
def isolated_client(isolated_engine, tmp_path):
    sf = sessionmaker(
        bind=isolated_engine,
        autoflush=False,
        expire_on_commit=False,
    )

    queue = InMemoryQueue()
    storage = FakeStorage(tmp_path / "objects")
    storage.root.mkdir(parents=True, exist_ok=True)

    orig_q = documents.queue
    orig_s = documents.storage
    documents.queue = queue
    documents.storage = storage

    def override() -> Iterator[Session]:
        s = sf()
        try:
            yield s
        finally:
            s.close()

    app.dependency_overrides[get_db_session] = override

    with TestClient(app) as c:
        yield c, queue, isolated_engine

    app.dependency_overrides.clear()
    documents.queue = orig_q
    documents.storage = orig_s


def _upload_and_process(client, queue, engine, job_id, count=1):
    inference = FakeInferenceClient()
    for i in range(count):
        client.post(
            f"/jobs/{job_id}/documents",
            files={
                "file": (
                    f"card{i}.jpg",
                    BytesIO(f"img-{i}".encode()),
                    "image/jpeg",
                )
            },
        )

    worker = Worker(
        queue=queue,
        processor=DocumentProcessor(
            engine=engine,
            object_reader=FakeObjectReader(),
            inference_client=inference,
        ),
    )
    for _ in range(count):
        worker.process_one()


def test_export_empty_job(isolated_client):
    client, _, _ = isolated_client

    job = client.post(
        "/jobs", json={"total_documents": 1}
    ).json()

    response = client.get(
        f"/jobs/{job['job_id']}/export/xlsx"
    )

    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]

    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    # Header row only.
    assert ws.max_row == 1
    assert ws.cell(row=1, column=1).value == "First Name"
    assert ws.cell(row=1, column=7).value == "Email Address"


def test_export_one_lead(isolated_client):
    client, queue, engine = isolated_client

    job = client.post(
        "/jobs", json={"total_documents": 1}
    ).json()
    job_id = job["job_id"]

    _upload_and_process(client, queue, engine, job_id, 1)

    response = client.get(f"/jobs/{job_id}/export/xlsx")

    assert response.status_code == 200

    wb = load_workbook(BytesIO(response.content))
    ws = wb.active

    assert ws.max_row == 2
    assert ws.cell(row=2, column=1).value == "First1"
    assert ws.cell(row=2, column=2).value == "Last1"
    assert ws.cell(row=2, column=7).value == "user1@test.com"


def test_export_multiple_leads(isolated_client):
    client, queue, engine = isolated_client

    job = client.post(
        "/jobs", json={"total_documents": 3}
    ).json()
    job_id = job["job_id"]

    _upload_and_process(client, queue, engine, job_id, 3)

    response = client.get(f"/jobs/{job_id}/export/xlsx")

    wb = load_workbook(BytesIO(response.content))
    ws = wb.active

    assert ws.max_row == 4  # header + 3 data rows


def test_export_correct_headers(isolated_client):
    client, _, _ = isolated_client

    job = client.post(
        "/jobs", json={"total_documents": 1}
    ).json()

    response = client.get(
        f"/jobs/{job['job_id']}/export/xlsx"
    )

    wb = load_workbook(BytesIO(response.content))
    ws = wb.active

    expected = [
        "First Name",
        "Last Name",
        "Position / Job Title",
        "Company",
        "Location",
        "Phone Number",
        "Email Address",
    ]

    for i, header in enumerate(expected, start=1):
        assert ws.cell(row=1, column=i).value == header


def test_export_unknown_job(isolated_client):
    client, _, _ = isolated_client

    response = client.get(
        "/jobs/00000000-0000-0000-0000-000000000000/export/xlsx"
    )

    assert response.status_code == 404


def test_export_content_disposition(isolated_client):
    client, _, _ = isolated_client

    job = client.post(
        "/jobs", json={"total_documents": 1}
    ).json()
    job_id = job["job_id"]

    response = client.get(f"/jobs/{job_id}/export/xlsx")

    cd = response.headers["content-disposition"]
    assert "attachment" in cd
    assert job_id in cd
    assert ".xlsx" in cd
