"""Excel export service for generating lead spreadsheets."""

from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from packages.common.models import LeadModel
from packages.common.repositories import LeadRepository


LEAD_COLUMNS = [
    ("First Name", "first_name"),
    ("Last Name", "last_name"),
    ("Position / Job Title", "job_title"),
    ("Company", "company"),
    ("Location", "location"),
    ("Phone Number", "phone_number"),
    ("Email Address", "email_address"),
]


def generate_leads_xlsx(
    *,
    session: Session,
    job_id: UUID,
) -> bytes:
    """Generate an XLSX file containing all leads for a job.

    Returns the file content as bytes.
    """
    repo = LeadRepository(session)
    leads: list[LeadModel] = repo.get_by_job_id(job_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    # Header row with styling.
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(
        start_color="2B579A",
        end_color="2B579A",
        fill_type="solid",
    )
    header_alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    for col_idx, (header_name, _) in enumerate(
        LEAD_COLUMNS, start=1
    ):
        cell = ws.cell(row=1, column=col_idx, value=header_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Data rows.
    for row_idx, lead in enumerate(leads, start=2):
        for col_idx, (_, attr_name) in enumerate(
            LEAD_COLUMNS, start=1
        ):
            value = getattr(lead, attr_name, None) or ""
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Auto-size columns.
    for col_idx, (header_name, _) in enumerate(
        LEAD_COLUMNS, start=1
    ):
        max_length = len(header_name)

        for row in range(2, len(leads) + 2):
            cell_value = str(
                ws.cell(row=row, column=col_idx).value or ""
            )
            max_length = max(max_length, len(cell_value))

        adjusted_width = min(max_length + 4, 50)
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[col_letter].width = adjusted_width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return buffer.read()
